from pathlib import Path
from typing import List, Union
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from backend.core.logger.logger import get_logger
from backend.core.utils.regex.regex_utils import RegexUtils

logger = get_logger()

class excel_generator:
    def __init__(self, filepath: Union[str, Path] = None):
        self.filepath = Path(filepath) if filepath else None
        if self.filepath and self.filepath.exists():
            self.wb = load_workbook(self.filepath)
            self.ws = self.wb.active
            logger.info(f"加载现有文件：{self.filepath}")
        else:
            self.wb = Workbook()
            self.ws = self.wb.active
            logger.info("创建新的工作簿。")

    def set_headers(self, headers: List[str], style=True):
        for col, header in enumerate(headers, 1):
            cell = self.ws.cell(row=1, column=col, value=header)
            if style:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill(start_color="4472C4", fill_type='solid')

    def add_data_object_row(self, obj: object, attributes: List[str]):
        next_row = self.ws.max_row + 1
        for col_idx, attr in enumerate(attributes, 1):
            value = getattr(obj, attr, None)
            self.ws.cell(row=next_row, column=col_idx, value=value)

    def save(self, filepath: Union[str, Path] = None, create_dirs: bool = True):
        filepath = Path(filepath or self.filepath)
        if create_dirs and not filepath.parent.exists():
            filepath.parent.mkdir(parents=True, exist_ok=True)
        self.wb.save(filepath)
        logger.info(f"文件已保存至: {filepath.absolute()}")
